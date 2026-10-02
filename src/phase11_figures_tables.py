"""
Phase 11 — manuscript figures and tables, built only from saved results files (no refitting).
Re-run after any change in results; every number in the tables comes from the files below.

Figures (manuscript/figures/, PNG 600 dpi + PDF vector; widths 174 mm double / 85 mm single column):
  Fig1_flow            cohort flow diagram                      results/phase9_v2/cohort_flow.csv
  Fig2_calibration     reliability diagrams, deciles of risk    {tag}_reliability.csv
  Fig3_time_course     (a) event onset within the anesthetic    phase4 events + phase5 features
                       (b) pooled vs within-phase AUROC         {tag}_diagnostic_metrics.csv, {tag}_phase_stratified.csv
  Fig4_reference       AUROC of Model A/B vs timing-only and    run_comparisons.csv (set refs)
                       anesthetic-course models, both periods
  Fig5_shap            SHAP summary, Models A and B, 5 min      shap_{tag}_{A,B}_5min.parquet
  FigS1_decision       decision curves                          {tag}_decision_curve.csv
Tables (manuscript/tables/CPI_tables.docx + one CSV per table):
  Table 1  characteristics                table1.csv
  Table 2  discrimination & calibration   {tag}_diagnostic_metrics.csv
  Table 3  operating points               {tag}_diagnostic_metrics.csv
  Table 4  time course                    phase_stratified, run_comparisons (refs, ablations)
  S4 selection bias; S5 sensitivity analyses (Tests 3, 4, 6); S6 subgroups; S7 matched sampling;
  S8 repeat patients; S9 effective sample size; S10 latency; S11 surgery-period discrimination
Missing inputs are skipped with a message; everything else is still produced.
Requires matplotlib, pandas, pyarrow, python-docx (pip install python-docx).
Usage: python src/phase11_figures_tables.py
"""
import os
import sys
import re
import textwrap

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cpi_config import RESULTS_DIR, PROCESSED_DIR, PROJECT_ROOT, RANDOM_SEED

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

R9 = os.path.join(RESULTS_DIR, 'phase9_v2')
TAG = 'sqi_primary_w70'
FIG_DIR = os.path.join(PROJECT_ROOT, 'manuscript', 'figures')
TAB_DIR = os.path.join(PROJECT_ROOT, 'manuscript', 'tables')
H = ['5min', '10min', '15min']
HLAB = {'5min': '5 min', '10min': '10 min', '15min': '15 min'}
MM = 1 / 25.4
W2, W1 = 174 * MM, 85 * MM

# Validated categorical palette, first three slots (all-pairs CVD-safe); grey for references.
COL = {'A': '#2a78d6', 'B': '#eb6834', 'baseline': '#1baf7a', 'ref': '#8a8984', 'ref2': '#52514e'}
MARK = {'A': 'o', 'B': 's', 'baseline': '^'}
MLAB = {'A': 'Model A', 'B': 'Model B', 'baseline': 'Baseline'}
INK, INK2, GRID = '#0b0b0b', '#52514e', '#e4e3df'
FEAT = {'elapsed_time_s': 'Elapsed time', 'PPF20_CE_mean': 'Propofol Ce', 'RFTN20_CE_mean': 'Remifentanil Ce',
        'tiva_propofol': 'Propofol TCI running', 'preop_hb': 'Preoperative Hb', 'age': 'Age', 'bmi': 'BMI',
        'asa': 'ASA class', 'HR_dfa_alpha1': 'HR DFA α1', 'HR_dfa_alpha2': 'HR DFA α2', 'HR_sampen': 'HR sample entropy'}
for tr, nm in (('HR', 'HR'), ('ART_MBP', 'MAP'), ('PLETH_SPO2', 'SpO$_2$'), ('ETCO2', 'ETCO$_2$')):
    for st, sn in (('mean', 'mean'), ('slope', 'slope'), ('var', 'variance')):
        FEAT[f'{tr}_{st}'] = f'{nm} {sn}'

plt.rcParams.update({
    'font.family': 'sans-serif', 'font.sans-serif': ['Arial', 'Helvetica', 'DejaVu Sans'],
    'font.size': 7, 'axes.titlesize': 7.5, 'axes.labelsize': 7, 'xtick.labelsize': 6.5, 'ytick.labelsize': 6.5,
    'legend.fontsize': 6.5, 'axes.edgecolor': INK2, 'axes.labelcolor': INK, 'xtick.color': INK2,
    'ytick.color': INK2, 'axes.linewidth': 0.6, 'xtick.major.width': 0.6, 'ytick.major.width': 0.6,
    'axes.spines.top': False, 'axes.spines.right': False, 'axes.grid': True, 'grid.color': GRID,
    'grid.linewidth': 0.5, 'axes.axisbelow': True, 'legend.frameon': False, 'pdf.fonttype': 42,
    'savefig.dpi': 600, 'figure.dpi': 150, 'mathtext.default': 'regular'})


def path(name):
    return os.path.join(R9, name)


def read(name, **kw):
    p = path(name)
    if not os.path.exists(p):
        print(f"  skipped: {name} not found")
        return None
    return pd.read_parquet(p, **kw) if p.endswith('.parquet') else pd.read_csv(p, **kw)


def save(fig, name):
    for ext in ('png', 'pdf'):
        fig.savefig(os.path.join(FIG_DIR, f'{name}.{ext}'), bbox_inches='tight')
    plt.close(fig)
    print(f"  {name}.png / .pdf")


def panel_label(ax, s):
    ax.text(-0.14, 1.04, s, transform=ax.transAxes, fontweight='bold', fontsize=8.5, va='bottom', color=INK)


def f3(x):
    return '' if pd.isna(x) else f'{x:.3f}'


def neg0(txt):
    """'-0.000' and '+0.000' -> '0.000' (rounded-away signs)."""
    return re.sub(r'(?<![\d.])[-+](0\.0+)(?![\d])', r'\1', txt) if isinstance(txt, str) else txt


def minus(txt):
    """Hyphen used as a minus sign -> true minus (Word tables only)."""
    return re.sub(r'(?<![\w.])-(?=\d)', '−', txt) if isinstance(txt, str) else txt


def ci(x, lo, hi, d=3):
    return '' if pd.isna(x) else f'{x:.{d}f} ({lo:.{d}f}–{hi:.{d}f})'


# ================================================================== figures
# Shorter wording for the flow diagram; the full reasons stay in cohort_flow.csv. Counts come from the file.
FLOW_SHORT = [
    (r'any involving HR/SpO2/ETCO2', 'HR, SpO2 or ETCO2'),
    (r'track list could not be downloaded (\d+) \(caseid \d+\)', r'track list unavailable \1'),
    (r'case-level missingness above cohort 95th percentile \(HR > ([\d.]+)% missing: (\d+); MAP > ([\d.]+)% missing: (\d+); '
     r'BIS final 45 min > ([\d.]+)% missing: (\d+); cases above >= 2 thresholds (\d+)\)',
     r'missingness above the cohort 95th percentile: HR > \1% (\2), MAP > \3% (\4), BIS in the final 45 min > \5% (\6); '
     r'\7 cases above ≥ 2 thresholds'),
    (r'no window with >= 70% completeness and an observable label \(case (\d+):.*?; case (\d+):.*\)$',
     r'no window ≥ 70% complete with an observable label (case \1: none passes the completeness rule; case \2: all labels censored)'),
    (r'>=', '≥'), (r'(?<![\d.,])(\d{4,})(?![\d.%])', lambda m: f'{int(m.group(1)):,}')]


def short_reason(txt):
    txt = str(txt)
    for pat, sub in FLOW_SHORT:
        txt = re.sub(pat, sub, txt)
    return txt


def fig1_flow():
    flow = read('cohort_flow.csv')
    if flow is None:
        return
    n = len(flow)
    fig, ax = plt.subplots(figsize=(W2, 0.66 * n + 0.2))
    ax.set_xlim(0, 10); ax.set_ylim(-0.2, n * 1.0); ax.axis('off')
    for i, r in flow.reset_index(drop=True).iterrows():
        y = n - 1 - i
        ax.add_patch(plt.Rectangle((0.2, y + 0.08), 3.6, 0.68, fill=False, lw=0.7, ec=INK))
        lab = textwrap.fill(str(r['step']).replace('>=', '≥'), 40)
        ax.text(2.0, y + 0.42, f"{lab}\nn = {int(r['n_remaining']):,}", ha='center', va='center', fontsize=6.5, color=INK)
        if i < n - 1:
            ax.annotate('', xy=(2.0, y - 0.24), xytext=(2.0, y + 0.08), arrowprops=dict(arrowstyle='-|>', lw=0.7, color=INK, mutation_scale=7))
        if i > 0 and pd.notna(r['n_excluded']) and str(r['n_excluded']).strip() not in ('', '0'):
            txt = textwrap.fill(f"Excluded, n = {int(float(r['n_excluded'])):,}: {short_reason(r['reason'])}", 84)
            ax.plot([2.0, 4.3], [y + 0.92, y + 0.92], lw=0.6, color=INK2)
            ax.text(4.35, y + 0.92, txt, ha='left', va='center', fontsize=5.6, color=INK2)
    save(fig, 'Fig1_flow')


def fig2_calibration():
    rel = read(f'{TAG}_reliability.csv')
    if rel is None:
        return
    fig, axes = plt.subplots(1, 3, figsize=(W2, W2 / 3 + 0.35))
    for k, (ax, h) in enumerate(zip(axes, H)):
        d = rel[rel['horizon'] == h]
        v = pd.concat([d['mean_pred'], d['obs_rate']])
        v = v[v > 0]
        lo, top = v.min() / 1.6, v.max() * 1.6
        ax.plot([lo, top], [lo, top], color=INK2, lw=0.7, ls=(0, (3, 2)), zorder=1)
        for m in ('A', 'B', 'baseline'):
            g = d[d['model'] == m].sort_values('mean_pred')
            ax.plot(g['mean_pred'], g['obs_rate'], color=COL[m], lw=1.2, marker=MARK[m], ms=3.2,
                    mec='white', mew=0.5, label=MLAB[m], zorder=3)
        ax.set_xscale('log'); ax.set_yscale('log')
        ax.set_xlim(lo, top); ax.set_ylim(lo, top); ax.set_aspect('equal')
        ax.grid(which='minor', visible=False)
        ax.set_title(HLAB[h] + ' horizon', color=INK)
        ax.set_xlabel('Mean predicted risk, by decile (log)')
        if k == 0:
            ax.set_ylabel('Observed event rate (log)')
        panel_label(ax, 'abc'[k])
    axes[0].legend(loc='upper left', handlelength=1.6)
    fig.tight_layout(w_pad=1.2)
    save(fig, 'Fig2_calibration')


def event_onsets():
    """Relative onset (0 = anesthesia start, 1 = end) of every primary event within anesthesia."""
    try:
        feats = pd.read_parquet(os.path.join(PROCESSED_DIR, 'phase5_features_v2.parquet'),
                                columns=['caseid', 't', 't_case_s', 'elapsed_time_s', 'aneend_case_s'])
        cohort = set(pd.read_csv(os.path.join(RESULTS_DIR, 'phase4_final_cohort_caseids.csv'))['caseid'])
        ev = pd.read_csv(os.path.join(RESULTS_DIR, 'phase4_events_sqi_primary.csv'))
    except FileNotFoundError as e:
        print(f"  skipped event onsets: {e}")
        return None
    pc = feats.groupby('caseid').first()
    pc['anestart'] = pc['t_case_s'] - pc['elapsed_time_s']
    pc['offset'] = pc['t_case_s'] - pc['t']
    pc['dur'] = pc['aneend_case_s'] - pc['anestart']
    ev = ev[ev['caseid'].isin(cohort)].join(pc[['anestart', 'offset', 'dur']], on='caseid', how='inner')
    rel = (pd.to_timedelta(ev['start']).dt.total_seconds() + ev['offset'] - ev['anestart']) / ev['dur']
    return rel[(rel >= 0) & (rel <= 1)].to_numpy()


def fig3_time_course():
    rel = event_onsets()
    dm, ps = read(f'{TAG}_diagnostic_metrics.csv'), read(f'{TAG}_phase_stratified.csv')
    if rel is None and (dm is None or ps is None):
        return
    fig, axes = plt.subplots(1, 2, figsize=(W2, 2.4), gridspec_kw={'width_ratios': [1, 1.1]})
    ax = axes[0]
    if rel is not None:
        bins = np.linspace(0, 1, 21)
        for lo, hi, lab in ((0, 0.4, 'early'), (0.4, 0.8, 'mid'), (0.8, 1.0, 'late')):
            ax.axvspan(lo, hi, color='#f0efec' if lab != 'mid' else '#f7f6f3', lw=0, zorder=0)
            share = ((rel >= lo) & (rel < hi if hi < 1 else rel <= hi)).mean() * 100
            ax.text((lo + hi) / 2, 1.0, f'{lab}\n{share:.0f}%', transform=ax.get_xaxis_transform(),
                    ha='center', va='top', fontsize=6.3, color=INK2)
        ax.hist(rel, bins=bins, color=COL['A'], edgecolor='white', linewidth=0.6, zorder=2)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, ax.get_ylim()[1] * 1.22)
        ax.set_xlabel('Event onset, fraction of the anesthetic\n(0 = anesthesia start, 1 = end)')
        ax.set_ylabel(f'Events (n = {len(rel):,})')
        ax.grid(axis='x', visible=False)
    panel_label(ax, 'a')
    ax = axes[1]
    if dm is not None and ps is not None:
        h = '5min'
        xs = np.array([0, 1.4, 2.4, 3.4])
        for i, m in enumerate(('A', 'B', 'baseline')):
            p = dm[(dm['model'] == m) & (dm['horizon'] == h)].iloc[0]
            g = ps[(ps['model'] == m) & (ps['horizon'] == h)].set_index('phase').reindex(['early', 'mid', 'late'])
            v = np.r_[p['auroc'], g['auroc'].to_numpy()]
            lo = np.r_[p['auroc_ci_lo'], g['auroc_ci_lo'].to_numpy()]
            hi = np.r_[p['auroc_ci_hi'], g['auroc_ci_hi'].to_numpy()]
            x = xs + (i - 1) * 0.16
            ax.errorbar(x, v, yerr=[v - lo, hi - v], fmt='none', ecolor=COL[m], elinewidth=0.8, capsize=0, zorder=2)
            ax.plot(x[1:], v[1:], color=COL[m], lw=1.0, zorder=2)
            ax.plot(x, v, ls='none', marker=MARK[m], ms=3.8, color=COL[m], mec='white', mew=0.5, label=MLAB[m], zorder=3)
        ax.axhline(0.5, color=INK2, lw=0.6, ls=(0, (3, 2)))
        ax.axvline(0.7, color=GRID, lw=0.8)
        ax.set_xticks(xs, ['Pooled', 'Early', 'Mid', 'Late']); ax.set_xlim(-0.5, 3.8)
        ax.grid(axis='x', visible=False)
        ax.set_ylabel('AUROC, 5-minute horizon (95% CI)'); ax.set_ylim(0.4, 0.95)
        ax.text(2.4, -0.2, 'within phase of the anesthetic', transform=ax.get_xaxis_transform(), ha='center', va='top', fontsize=6.3, color=INK2)
        ax.legend(loc='lower center', bbox_to_anchor=(0.5, 1.0), ncol=3, handlelength=1.2, borderaxespad=0.3)
    panel_label(ax, 'b')
    fig.tight_layout(w_pad=2)
    save(fig, 'Fig3_time_course')


def fig4_reference():
    rc = read('run_comparisons.csv')
    if rc is None:
        return
    refs = rc[rc['set'] == 'refs']
    if refs.empty:
        print("  skipped Fig4: no 'refs' rows"); return
    fig, axes = plt.subplots(1, 2, figsize=(W2, 2.3), sharey=True)
    for k, (ax, per, title) in enumerate(zip(axes, ('anesthesia', 'surgery'), ('Anesthesia period (primary)', 'Surgery period (sensitivity analysis)'))):
        a = refs[refs['comparison'] == f'Model A vs anesthetic course ({per})'].set_index('horizon').reindex(H)
        b = refs[refs['comparison'] == f'Model B vs timing only ({per})'].set_index('horizon').reindex(H)
        x = np.arange(3)
        series = [('Model A', a['auroc_1'], COL['A'], '-', 'o'), ('Elapsed time + drug features', a['auroc_2'], COL['ref2'], (0, (4, 2)), 'D'),
                  ('Model B', b['auroc_1'], COL['B'], '-', 's'), ('Elapsed time only', b['auroc_2'], COL['ref'], (0, (1.5, 1.5)), 'v')]
        for lab, v, c, ls, mk in series:
            ax.plot(x, v.to_numpy(), color=c, ls=ls, lw=1.3, marker=mk, ms=3.6, mec='white', mew=0.5, label=lab)
        ends = sorted((v.iloc[-1], v.iloc[-1]) for _, v, *_ in series)   # (label y, true y), spread >= 0.012
        ys = [e[0] for e in ends]
        for i in range(1, len(ys)):
            ys[i] = max(ys[i], ys[i - 1] + 0.012)
        for yl, (_, yv) in zip(ys, ends):
            ax.annotate(f'{yv:.2f}', (x[-1], yv), xytext=(x[-1] + 0.1, yl), textcoords='data', va='center', fontsize=6, color=INK2)
        ax.set_xticks(x, [HLAB[h] for h in H]); ax.set_xlim(-0.2, 2.45)
        ax.set_title(title, color=INK); ax.set_xlabel('Prediction horizon')
        if k == 0:
            ax.set_ylabel('AUROC')
        panel_label(ax, 'ab'[k])
    lo = min(refs[['auroc_1', 'auroc_2']].min().min() - 0.02, 0.6)
    axes[0].set_ylim(np.floor(lo * 20) / 20, 0.9)
    fig.tight_layout(w_pad=1.5)
    h_, l_ = axes[0].get_legend_handles_labels()
    fig.legend(h_, l_, loc='upper center', bbox_to_anchor=(0.5, 0.0), ncol=4, handlelength=2.6)
    save(fig, 'Fig4_reference_models')


def beeswarm(ax, df, top=10, rng=None):
    feats = [c[5:] for c in df.columns if c.startswith('shap_')]
    imp = pd.Series({f: df[f'shap_{f}'].abs().mean() for f in feats}).sort_values(ascending=False).iloc[:top][::-1]
    cmap = plt.get_cmap('coolwarm')
    for i, f in enumerate(imp.index):
        v, x = df[f'shap_{f}'].to_numpy(), df[f].to_numpy(dtype=float)
        ok = np.isfinite(x)
        lo, hi = (np.nanpercentile(x, 5), np.nanpercentile(x, 95)) if ok.any() else (0, 1)
        c = np.clip((x - lo) / (hi - lo if hi > lo else 1), 0, 1)
        yj = i + rng.uniform(-0.32, 0.32, len(v))
        ax.scatter(v[~ok], yj[~ok], s=1.2, c='#b4b3ad', lw=0, alpha=0.6, rasterized=True)
        sc = ax.scatter(v[ok], yj[ok], s=1.2, c=c[ok], cmap=cmap, vmin=0, vmax=1, lw=0, alpha=0.7, rasterized=True)
    ax.axvline(0, color=INK2, lw=0.6)
    ax.set_yticks(range(len(imp)), [FEAT.get(f, f) for f in imp.index])
    ax.grid(axis='y', visible=False)
    ax.set_xlabel('SHAP value (log-odds)')
    return sc


def fig5_shap():
    rng = np.random.RandomState(RANDOM_SEED)
    fig = plt.figure(figsize=(W2, 2.9))
    gs = fig.add_gridspec(1, 2, wspace=0.75, right=0.9)
    axes = [fig.add_subplot(gs[0, 0]), fig.add_subplot(gs[0, 1])]
    cax = axes[1].inset_axes([1.05, 0.0, 0.035, 1.0])
    sc = None
    for k, (ax, m) in enumerate(zip(axes, ('A', 'B'))):
        df = read(f'shap_{TAG}_{m}_5min.parquet')
        if df is None:
            ax.axis('off'); continue
        sc = beeswarm(ax, df, rng=rng)
        ax.set_title(MLAB[m], color=INK)
        ax.text(-0.62, 1.04, 'ab'[k], transform=ax.transAxes, fontweight='bold', fontsize=8.5, va='bottom', color=INK)
    if sc is not None:
        cb = fig.colorbar(sc, cax=cax, ticks=[0, 1])
        cb.ax.set_yticklabels(['low', 'high']); cb.set_label('Feature value', fontsize=6.5)
        cb.outline.set_linewidth(0.4)
        fig.text(0.5, -0.07, '5-minute horizon; ten features with the largest mean |SHAP|. Grey points: feature missing in that window.',
                 ha='center', fontsize=6, color=INK2)
    else:
        cax.axis('off')
    save(fig, 'Fig5_shap')


def figS1_decision():
    dca = read(f'{TAG}_decision_curve.csv')
    if dca is None:
        return
    fig, axes = plt.subplots(1, 3, figsize=(W2, 2.2), sharey=False)
    for k, (ax, h) in enumerate(zip(axes, H)):
        d = dca[dca['horizon'] == h]
        d0 = d[d['model'] == d['model'].iloc[0]]
        ax.plot(d0['threshold'], d0['net_benefit_treat_all'], color=COL['ref'], lw=1, ls=(0, (3, 2)), label='Alarm always')
        ax.axhline(0, color=INK2, lw=0.8, label='Alarm never')
        for m in ('A', 'B', 'baseline'):
            g = d[d['model'] == m]
            ax.plot(g['threshold'], g['net_benefit_model'], color=COL[m], lw=1.3, label=MLAB[m])
        ymax = d['net_benefit_model'].max() * 1.15
        ax.set_ylim(-ymax * 0.25, ymax); ax.set_xlim(0, d['threshold'].max())
        ax.set_title(HLAB[h] + ' horizon', color=INK); ax.set_xlabel('Threshold probability')
        if k == 0:
            ax.set_ylabel('Net benefit')
        panel_label(ax, 'abc'[k])
    axes[0].legend(loc='upper right')
    fig.tight_layout(w_pad=1.2)
    save(fig, 'FigS1_decision_curves')


# ================================================================== tables
TABLES = []   # (title, DataFrame, footnote)


def add(title, df, note=''):
    df = df.fillna('').astype(str).apply(lambda c: c.map(neg0))
    TABLES.append((title, df, note))
    fn = title.split('.')[0].replace(' ', '_')
    df.to_csv(os.path.join(TAB_DIR, f'{fn}.csv'), index=False)


VAR_NAMES = {'Patients (unique subjectid), n': 'Patients, n', '   F': '   Female', '   M': '   Male', '   other': '   All other types'}


def nice_var(v):
    v = VAR_NAMES.get(v, v)
    return v.replace('kg/m2', 'kg/m²').replace('BIS>=60', 'BIS ≥ 60').replace('>=', '≥')


def collapse_binary(t, col='variable'):
    """'Emergency surgery, n (%)' followed by a lone '   1' row -> one row."""
    t = t.reset_index(drop=True)
    drop = []
    for i in range(len(t) - 1):
        nxt = t.loc[i + 1, col]
        after = t.loc[i + 2, col] if i + 2 < len(t) else ''
        if nxt.strip() == '1' and nxt.startswith(' ') and not str(after).startswith(' '):
            vals = t.loc[i + 1].drop(col)
            t.loc[i, vals.index] = vals.values
            drop.append(i + 1)
    return t.drop(index=drop).reset_index(drop=True)


def table1():
    t = read('table1.csv', keep_default_na=False)
    if t is None:
        return
    t = t[~t['variable'].str.strip().isin(['V'])]
    t['variable'] = t['variable'].map(nice_var)
    t = t.rename(columns={'variable': 'Characteristic', 'At least one event': '≥1 event', 'missing': 'Missing, n (%)'})
    add('Table 1. Characteristics of the included cases, overall and by occurrence of at least one sustained BIS ≥ 60 event during anesthesia',
        t, 'Values are mean (SD), median [IQR] or n (%). SMD, standardized mean difference between cases with and without an event. '
           'ASA class 6 recoded as missing. SMD, cases with an event minus cases without. Event rows define the groups and have no SMD. '
           'Two of the 2,673 cases contributed no analysable window (Figure 1).')


def table2_3():
    d = read(f'{TAG}_diagnostic_metrics.csv')
    if d is None:
        return
    d = d.set_index(['horizon', 'model']).reindex([(h, m) for h in H for m in ('A', 'B', 'baseline')]).reset_index()
    t2 = pd.DataFrame({'Horizon': d['horizon'].map(HLAB), 'Model': d['model'].map(MLAB),
                       'AUROC (95% CI)': [ci(r.auroc, r.auroc_ci_lo, r.auroc_ci_hi) for r in d.itertuples()],
                       'AUPRC (95% CI)': [ci(r.auprc, r.auprc_ci_lo, r.auprc_ci_hi) for r in d.itertuples()],
                       'Brier (95% CI)': [ci(r.brier, r.brier_ci_lo, r.brier_ci_hi, 4) for r in d.itertuples()],
                       'ECE': d['ece'].map(lambda x: '<0.0001' if x < 0.00005 else f'{x:.4f}'),
                       'Calibration intercept': d['calibration_intercept'].map(lambda x: f'{x:.3f}'),
                       'Calibration slope': d['calibration_slope'].map(lambda x: f'{x:.2f}')})
    prev = d.groupby('horizon')['prevalence'].first() if 'prevalence' in d else None
    note = 'Pooled out-of-fold predictions after nested calibration; 95% CIs from 1,000 case-level bootstrap resamples. ECE, expected calibration error (10 equal-width bins).'
    if prev is not None:
        note += ' Prevalence: ' + ', '.join(f'{HLAB[h]} {prev[h]*100:.2f}%' for h in H) + '.'
    add('Table 2. Discrimination and calibration by model and prediction horizon', t2, note)
    t3 = pd.DataFrame({'Horizon': d['horizon'].map(HLAB), 'Model': d['model'].map(MLAB),
                       'Sensitivity': [ci(r.sensitivity, r.sensitivity_ci_lo, r.sensitivity_ci_hi) for r in d.itertuples()],
                       'Specificity': [ci(r.specificity, r.specificity_ci_lo, r.specificity_ci_hi) for r in d.itertuples()],
                       'PPV': [ci(r.ppv, r.ppv_ci_lo, r.ppv_ci_hi) for r in d.itertuples()],
                       'NPV': [ci(r.npv, r.npv_ci_lo, r.npv_ci_hi) for r in d.itertuples()],
                       'Alarms per hour': [ci(r.alarms_per_hour, r.alarms_per_hour_ci_lo, r.alarms_per_hour_ci_hi, 1) for r in d.itertuples()],
                       'Events anticipated': [ci(r.event_sensitivity, r.event_sensitivity_ci_lo, r.event_sensitivity_ci_hi) for r in d.itertuples()],
                       'Median lead time, min': d['lead_time_median_min'].map(lambda x: f'{x:.1f}')})
    add('Table 3. Operating characteristics at the cross-fitted Youden threshold', t3,
        'Each outer fold\'s threshold was chosen on the other four folds. Alarms per hour of eligible monitoring time. '
        'Events anticipated: share of events with at least one alarm within the horizon before onset; lead time is capped by the horizon.')


def table4():
    ps, dm, rc = read(f'{TAG}_phase_stratified.csv'), read(f'{TAG}_diagnostic_metrics.csv'), read('run_comparisons.csv')
    rows = []
    if dm is not None:
        for m in ('A', 'B', 'baseline'):
            g = dm[dm['model'] == m].set_index('horizon')
            rows.append({'Analysis': f'{MLAB[m]}, pooled', **{HLAB[h]: f3(g.loc[h, 'auroc']) for h in H}})
    if ps is not None:
        for m in ('A', 'B', 'baseline'):
            g = ps[ps['model'] == m]
            rows.append({'Analysis': f'{MLAB[m]}, within phase (early / mid / late)',
                         **{HLAB[h]: ' / '.join(f3(g[(g['horizon'] == h) & (g['phase'] == p)]['auroc'].iloc[0]) for p in ('early', 'mid', 'late')) for h in H}})
        g = ps[ps['model'] == 'A']
        rows.append({'Analysis': 'Model A, events anticipated (early / mid / late)',
                     **{HLAB[h]: ' / '.join(f"{g[(g['horizon'] == h) & (g['phase'] == p)]['event_sensitivity'].iloc[0]*100:.0f}%" for p in ('early', 'mid', 'late')) for h in H}})
    if rc is not None:
        def val(comp, col):
            r = rc[rc['comparison'] == comp].set_index('horizon')
            return {HLAB[h]: f3(r.loc[h, col]) if h in r.index else '' for h in H}

        def diff(comp):
            r = rc[rc['comparison'] == comp].set_index('horizon')
            return {HLAB[h]: (f"{r.loc[h, 'auroc_diff']:+.3f} ({r.loc[h, 'auroc_diff_ci_lo']:+.3f} to {r.loc[h, 'auroc_diff_ci_hi']:+.3f})"
                              if h in r.index else '') for h in H}
        for per, lab in (('anesthesia', 'Anesthesia period'), ('surgery', 'Surgery period')):
            rows.append({'Analysis': f'{lab}: elapsed time only', **val(f'Model B vs timing only ({per})', 'auroc_2')})
            rows.append({'Analysis': f'{lab}: elapsed time + drug features', **val(f'Model A vs anesthetic course ({per})', 'auroc_2')})
            if per == 'surgery':
                rows.append({'Analysis': 'Surgery period: Model A', **val('Model A vs anesthetic course (surgery)', 'auroc_1')})
                rows.append({'Analysis': 'Surgery period: Model B', **val('Model B vs timing only (surgery)', 'auroc_1')})
            rows.append({'Analysis': f'{lab}: Model A − (elapsed time + drug features), ΔAUROC (95% CI)', **diff(f'Model A vs anesthetic course ({per})')})
            rows.append({'Analysis': f'{lab}: Model B − elapsed time only, ΔAUROC (95% CI)', **diff(f'Model B vs timing only ({per})')})
        if (rc['set'] == 'ablations').any():
            rows.append({'Analysis': 'Model A without drug features', **val('Model A vs Model A without drug features', 'auroc_2')})
            rows.append({'Analysis': 'Model A − Model A without drug features, ΔAUROC (95% CI)', **diff('Model A vs Model A without drug features')})
            rows.append({'Analysis': 'Model A without drug features − Model B, ΔAUROC (95% CI)', **diff('Model A without drug features vs Model B')})
            rows.append({'Analysis': 'Model A − Model A without elapsed time, ΔAUROC (95% CI)', **diff('Model A vs Model A without elapsed time')})
            rows.append({'Analysis': 'Model B − Model B without elapsed time, ΔAUROC (95% CI)', **diff('Model B vs Model B without elapsed time')})
    if rows:
        add('Table 4. The time course of the anesthetic: pooled and within-phase discrimination, reference models and ablations (AUROC)',
            pd.DataFrame(rows), 'Phases by relative time within the anesthetic: early < 0.4, mid 0.4–0.8, late > 0.8. Reference models and '
            'ablations refit with the primary configurations and folds; differences from paired case-level bootstrap (1,000 resamples).')


def supplementary():
    sb = read('selection_bias.csv', keep_default_na=False)
    if sb is not None:
        sb = sb[~sb['variable'].str.strip().isin(['V'])]
        sb = collapse_binary(sb)
        sb['variable'] = sb['variable'].map(nice_var)
        sb.columns = [c.strip() for c in sb.columns]
        sb = sb.rename(columns={'variable': 'Characteristic', 'Excluded, data availability (all)': 'Excluded, data availability: all',
                                'missing track(s)': 'of which missing track(s)', 'extreme missingness': 'of which extreme missingness'})
        add('Table S4. Characteristics of retained cases and of excluded cases, by reason for exclusion', sb,
            'Excluded by design: not general anesthesia, liver transplantation or donor hepatectomy, recording < 45 min. '
            'SMD compares retained cases with cases excluded for data availability. Duration SMDs are mean-based and distorted by extreme values; compare medians.')
    rc = read('run_comparisons.csv')
    if rc is not None:
        s = rc[rc['set'].isin(['test3', 'test4', 'test6'])]
        t = pd.DataFrame({'Comparison': s['comparison'], 'Model': s['model_1'].map(MLAB), 'Horizon': s['horizon'].map(HLAB),
                          'Common windows': s['n_windows_common'].map(lambda x: f'{int(x):,}'),
                          'AUROC primary': s['auroc_1'].map(f3), 'AUROC variant': s['auroc_2'].map(f3),
                          'ΔAUROC (95% CI)': [f'{r.auroc_diff:+.3f} ({r.auroc_diff_ci_lo:+.3f} to {r.auroc_diff_ci_hi:+.3f})' for r in s.itertuples()],
                          'p': s['auroc_p'].map(lambda x: '<0.001' if x <= 0.001 else f'{x:.3f}'),
                          'ΔAUPRC (95% CI)': [f'{r.auprc_diff:+.3f} ({r.auprc_diff_ci_lo:+.3f} to {r.auprc_diff_ci_hi:+.3f})' for r in s.itertuples()]})
        add('Table S5. Sensitivity analyses: label filtering (Tests 3–4) and window completeness (Test 6)', t,
            'Primary: SQI < 50 labels, windows with every core vital sign ≥ 70% present. Each variant compared on the windows it shares with the primary run, each with its own labels; paired case-level bootstrap.')
    sg = read(f'{TAG}_subgroups.csv')
    if sg is not None:
        t = pd.DataFrame({'Model': sg['model'].map(MLAB), 'Horizon': sg['horizon'].map(HLAB), 'Variable': sg['variable'],
                          'Group': sg['group'], 'Cases': sg['n_cases'], 'Positive windows': sg['n_positive'],
                          'AUROC (95% CI)': [ci(r.auroc, r.auroc_ci_lo, r.auroc_ci_hi) if pd.notna(getattr(r, 'auroc', np.nan)) else '' for r in sg.itertuples()],
                          'Observed/expected': sg['obs_exp_ratio'].map(lambda x: '' if pd.isna(x) else f'{x:.2f}'),
                          'Calibration slope': sg['calibration_slope'].map(lambda x: '' if pd.isna(x) else f'{x:.2f}')})
        add('Table S6. Discrimination and calibration by ASA class and age group', t,
            'Descriptive; no formal tests between subgroups. CIs from 500 case-level bootstrap resamples. The ASA IV–V group (24 cases) is too small to interpret.')
    ms = read('matched_sampling_results.csv')
    if ms is not None:
        dif = ms['model'].eq('A minus B')
        t = pd.DataFrame({'Horizon': ms['horizon'].map(HLAB), 'Model': ms['model'].map(lambda m: MLAB.get(m, 'Model A − Model B')),
                          'Positive / negative windows': [f'{int(a):,} / {int(b):,}' if pd.notna(a) else '' for a, b in zip(ms.get('n_pos', np.nan), ms.get('n_neg', np.nan))],
                          'AUROC (95% CI)': [(f'{r.auroc:+.3f} ({r.auroc_ci_lo:+.3f} to {r.auroc_ci_hi:+.3f})' if d else ci(r.auroc, r.auroc_ci_lo, r.auroc_ci_hi))
                                             for r, d in zip(ms.itertuples(), dif)],
                          'p (A vs B)': ms.get('auroc_p', pd.Series(np.nan, index=ms.index)).map(lambda x: '' if pd.isna(x) else ('<0.001' if x <= 0.001 else f'{x:.3f}'))})
        add('Table S7. Matched negative-sampling secondary analysis', t,
            'Negatives: up to two random windows per case, ≥ 20 min from any event and ≥ 10 min apart, shared across horizons; positives: all positive windows. AUPRC not shown (driven by the 62–82% positive fraction).')
    rp = read('repeat_patient_sensitivity.csv')
    if rp is not None:
        m = rp.copy()
        dif = m['model'].str.startswith('A minus')
        t = pd.DataFrame({'Horizon': m['horizon'].map(HLAB), 'Model': m['model'].map(lambda x: MLAB.get(x, 'Model A − Model B')),
                          'AUROC, all windows': [ci(r.auroc_all, r.auroc_all_lo, r.auroc_all_hi) for r in m.itertuples()],
                          'AUROC, restricted': [(f'{r.auroc_restricted:+.3f} ({r.auroc_restricted_lo:+.3f} to {r.auroc_restricted_hi:+.3f})' if d
                                                 else ci(r.auroc_restricted, r.auroc_restricted_lo, r.auroc_restricted_hi)) for r, d in zip(m.itertuples(), dif)],
                          'Change': m['auroc_change'].map(lambda x: '' if pd.isna(x) else f'{x:+.4f}')})
        add('Table S8. Repeat-patient sensitivity analysis', t,
            'Restricted: all 98 cases of the 43 patients with surgeries in different outer folds removed; bootstrap resamples patients.')
    en = read('effective_sample_size.csv')
    if en is not None:
        t = pd.DataFrame({'Horizon': en['horizon'].map(HLAB), 'Windows': en['n_windows'].map(lambda x: f'{int(x):,}'),
                          'Cases': en['n_cases'], 'Positive windows': en['n_pos_windows'].map(lambda x: f'{int(x):,}'),
                          'ICC': en['icc_label'].map(lambda x: f'{x:.3f}'), 'Design effect': en['de_unequal'].map(lambda x: f'{x:.1f}'),
                          'Effective windows': en['n_eff_windows'].map(lambda x: f'{x:,.0f}'),
                          'Effective positives': en['n_eff_pos'].map(lambda x: f'{x:,.0f}'),
                          'Effective positives per feature (A / B)': [f'{a:.0f} / {b:.0f}' for a, b in zip(en['eff_pos_per_feature_A'], en['eff_pos_per_feature_B'])],
                          'AUROC design effect (A / B)': [f'{a:.1f} / {b:.1f}' for a, b in zip(en['de_auroc_A'], en['de_auroc_B'])]})
        add('Table S9. Effective (clustering-adjusted) sample size', t,
            'Design effect with unequal-cluster correction (Kish; Eldridge et al.). AUROC design effect: variance under case-level vs window-level bootstrap.')
    lt = read('latency_summary.csv')
    if lt is not None:
        lt = lt.copy()
        lt['tier'] = lt['tier'].map({'workstation': 'All cores', 'single': 'Single thread, one core'}).fillna(lt['tier'])
        comp = {'hemodynamic': 'Hemodynamic features (4 signals)', 'dfa': 'HR DFA α1 and α2', 'sampen': 'HR sample entropy',
                'drug': 'Drug features', 'inference_A': 'Inference, Model A', 'inference_B': 'Inference, Model B',
                'total_A': 'Total, Model A', 'total_B': 'Total, Model B'}
        t = pd.DataFrame({'Configuration': lt['tier'], 'Component': lt['component'].map(lambda c: comp.get(c, c.replace('_', ' '))),
                          'Mean (SD), ms': [f'{a:.1f} ({b:.1f})' for a, b in zip(lt['mean_ms'], lt['sd_ms'])],
                          'Median, ms': lt['median_ms'].map(lambda x: f'{x:.1f}'), '95th percentile, ms': lt['p95_ms'].map(lambda x: f'{x:.1f}'),
                          'Maximum, ms': lt['max_ms'].map(lambda x: f'{x:.1f}')})
        add('Table S10. Per-window computation time from raw signals to a calibrated prediction', t,
            'Intel Core i7-1280P laptop processor, batch size 1, 1,000 windows after 50 warm-up windows. Inference: imputation of the '
            'clinical variables, gradient-boosting prediction and calibration map. Total, Model B excludes the drug features. One prediction is due every 60 s.')
    sd = read(f'{TAG}_surgery_diagnostic_metrics.csv')
    if sd is not None:
        sd = sd.set_index(['horizon', 'model']).reindex([(h, m) for h in H for m in ('A', 'B', 'baseline')]).dropna(how='all').reset_index()
        t = pd.DataFrame({'Horizon': sd['horizon'].map(HLAB), 'Model': sd['model'].map(MLAB),
                          'AUROC (95% CI)': [ci(r.auroc, r.auroc_ci_lo, r.auroc_ci_hi) for r in sd.itertuples()],
                          'AUPRC (95% CI)': [ci(r.auprc, r.auprc_ci_lo, r.auprc_ci_hi) for r in sd.itertuples()],
                          'Calibration slope': sd['calibration_slope'].map(lambda x: f'{x:.2f}'),
                          'Sensitivity': sd['sensitivity'].map(f3), 'Specificity': sd['specificity'].map(f3),
                          'Alarms per hour': sd['alarms_per_hour'].map(lambda x: f'{x:.1f}'),
                          'Events anticipated': sd['event_sensitivity'].map(f3)})
        add('Table S11. Surgery-period sensitivity analysis: discrimination, calibration and operating points', t,
            'Windows and events restricted to the surgical period; primary folds and model configurations reused.')


def write_docx():
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt, Mm
    except ImportError:
        print("  python-docx not installed: tables saved as CSV only (pip install python-docx for the Word file)")
        return
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Mm(15))
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(8)
    for i, (title, df, note) in enumerate(TABLES):
        if i:
            doc.add_page_break()
        p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(9)
        tb = doc.add_table(rows=1, cols=len(df.columns))
        tb.style = 'Table Grid'
        for j, c in enumerate(df.columns):
            cell = tb.rows[0].cells[j]; cell.text = ''
            run = cell.paragraphs[0].add_run(str(c)); run.bold = True; run.font.size = Pt(7.5)
        for _, row in df.iterrows():
            cells = tb.add_row().cells
            for j, v in enumerate(row.values):
                cells[j].text = ''
                run = cells[j].paragraphs[0].add_run(minus(str(v))); run.font.size = Pt(7.5)
        if note:
            q = doc.add_paragraph(); rn = q.add_run(note); rn.italic = True; rn.font.size = Pt(7)
    out = os.path.join(TAB_DIR, 'CPI_tables.docx')
    doc.save(out)
    print(f"  CPI_tables.docx ({len(TABLES)} tables)")


def main():
    os.makedirs(FIG_DIR, exist_ok=True); os.makedirs(TAB_DIR, exist_ok=True)
    print("Figures:")
    for f in (fig1_flow, fig2_calibration, fig3_time_course, fig4_reference, fig5_shap, figS1_decision):
        try:
            f()
        except Exception as e:
            print(f"  {f.__name__} FAILED: {type(e).__name__}: {e}")
    print("Tables:")
    for f in (table1, table2_3, table4, supplementary):
        try:
            f()
        except Exception as e:
            print(f"  {f.__name__} FAILED: {type(e).__name__}: {e}")
    for title, df, _ in TABLES:
        print(f"  {title.split('.')[0]}: {len(df)} rows")
    write_docx()
    print(f"\nFigures in {FIG_DIR}\nTables in {TAB_DIR}")


if __name__ == '__main__':
    main()